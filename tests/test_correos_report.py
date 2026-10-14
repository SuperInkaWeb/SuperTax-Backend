"""Exportación Excel (Fase 6): generación del .xlsx y mapeo de filas. Sin BD."""
import io
from datetime import UTC, date, datetime
from decimal import Decimal

from openpyxl import load_workbook

from src.modules.correos.application.report_service import _a_fila
from src.modules.correos.infrastructure.models import ExtractedRecordModel
from src.modules.correos.infrastructure.report_excel import (
    COLUMNAS,
    FilaReporte,
    generar_reporte,
)

_RUC = "20131312955"


def test_generar_reporte_cabecera_y_numeros():
    fila = FilaReporte(
        ruc_contribuyente=_RUC, nombre_contribuyente="INVERSIONES LIMA SAC",
        periodo="202603", importe_pagado="75.00", monto_alquiler="1500.00", moneda="PEN",
    )
    hoja = load_workbook(io.BytesIO(generar_reporte([fila]))).active
    cabecera = [c.value for c in hoja[1]]
    assert cabecera == [etiqueta for _, etiqueta in COLUMNAS]
    valores = {clave: hoja.cell(row=2, column=i).value for i, (clave, _) in enumerate(COLUMNAS, 1)}
    assert valores["ruc_contribuyente"] == _RUC
    # Los importes se escriben como número (para que la columna sume), no como texto.
    assert valores["importe_pagado"] == 75.0
    assert isinstance(valores["monto_alquiler"], (int, float))  # número, no cadena
    assert valores["monto_alquiler"] == 1500


def test_generar_reporte_vacio_solo_cabecera():
    hoja = load_workbook(io.BytesIO(generar_reporte([]))).active
    assert hoja.max_row == 1  # solo la cabecera


def test_a_fila_formatea_estado_fecha_e_importes():
    modelo = ExtractedRecordModel(
        id=1, job_id=2, attachment_id=3, ruc_contribuyente=_RUC,
        nombre_contribuyente="INVERSIONES LIMA SAC", ruc_inquilino=None, nombre_inquilino="",
        tipo_doc_inquilino="RUC", tipo_de_bien="PREDIO", periodo="202603",
        fecha_de_pago=date(2026, 4, 15), numero_de_operacion="987654321",
        monto_alquiler=Decimal("1500.00"), tributo_resultante=None,
        importe_pagado=Decimal("75.00"), intereses_moratorios=None, moneda="PEN",
        completitud="partial", estado_de_revision="pending",
        created_at=datetime.now(UTC),
    )
    fila = _a_fila(modelo, {"nombre_adjunto": "recibo.pdf"})
    assert fila.estado == "Parcial"
    assert fila.revision == "Pendiente"
    assert fila.fecha_de_pago == "15/04/2026"
    assert fila.importe_pagado == "75.00"
    assert fila.tributo_resultante == ""  # None → vacío
    assert fila.archivo_origen == "recibo.pdf"


def test_ruta_export_montada_antes_del_detalle():
    from src.main import create_app

    paths = create_app().openapi()["paths"]
    assert "/api/correos/records/export" in paths
    assert "get" in paths["/api/correos/records/export"]
