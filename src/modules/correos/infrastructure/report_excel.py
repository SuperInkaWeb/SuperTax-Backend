"""
Generadores del reporte de registros tributarios (constancias de arrendamiento):
Excel (.xlsx) y CSV, compartiendo las mismas columnas y formato de fila.

Excel con openpyxl en modo `write_only`: vuelca fila a fila en lugar de mantener la
hoja entera en memoria, para que un reporte grande no infle la RAM del proceso web.
Los importes se escriben como NÚMERO con formato de celda (no texto): un Excel donde
la columna de importes son cadenas no suma, y lo primero que hace un contador es
seleccionar la columna para ver el total.

CSV con BOM UTF-8 (`utf-8-sig`): sin él, Excel en Windows lee el fichero en la
codificación local y las tildes salen rotas (la queja número uno con exportaciones
en español).
"""
from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from typing import Final

# (clave, etiqueta) en el orden de las columnas del reporte.
COLUMNAS: Final[tuple[tuple[str, str], ...]] = (
    ("ruc_contribuyente", "RUC Arrendador"),
    ("nombre_contribuyente", "Nombre / Razón Social"),
    ("tipo_doc_inquilino", "Tipo Doc. Inquilino"),
    ("ruc_inquilino", "RUC Arrendatario"),
    ("nombre_inquilino", "Inquilino"),
    ("tipo_de_bien", "Tipo de Bien"),
    ("periodo", "Periodo"),
    ("monto_alquiler", "Monto Alquiler"),
    ("tributo_resultante", "Tributo Resultante"),
    ("importe_pagado", "Importe Pagado"),
    ("intereses_moratorios", "Intereses Moratorios"),
    ("moneda", "Moneda"),
    ("fecha_de_pago", "Fecha de Pago"),
    ("numero_de_operacion", "N° Operación"),
    ("estado", "Estado de Extracción"),
    ("revision", "Revisión"),
    ("archivo_origen", "Archivo"),
)

# Formato contable peruano: separador de miles y dos decimales.
_FORMATO_IMPORTE: Final = "#,##0.00"
_COLUMNAS_NUMERICAS: Final[frozenset[str]] = frozenset(
    {"monto_alquiler", "tributo_resultante", "importe_pagado", "intereses_moratorios"}
)
_ANCHOS: Final[dict[str, int]] = {
    "ruc_contribuyente": 14, "nombre_contribuyente": 38, "tipo_doc_inquilino": 16,
    "ruc_inquilino": 14, "nombre_inquilino": 38, "tipo_de_bien": 12, "periodo": 10,
    "monto_alquiler": 14, "tributo_resultante": 16, "importe_pagado": 14,
    "intereses_moratorios": 16, "moneda": 8, "fecha_de_pago": 13,
    "numero_de_operacion": 18, "estado": 18, "revision": 14, "archivo_origen": 34,
}


@dataclass(frozen=True, slots=True)
class FilaReporte:
    """Una fila del reporte, con todo ya convertido a texto (el formateo ocurre al
    construir la fila, no al escribir, para que no haya dos versiones del formato)."""

    ruc_contribuyente: str = ""
    nombre_contribuyente: str = ""
    tipo_doc_inquilino: str = ""
    ruc_inquilino: str = ""
    nombre_inquilino: str = ""
    tipo_de_bien: str = ""
    periodo: str = ""
    monto_alquiler: str = ""
    tributo_resultante: str = ""
    importe_pagado: str = ""
    intereses_moratorios: str = ""
    moneda: str = ""
    fecha_de_pago: str = ""
    numero_de_operacion: str = ""
    estado: str = ""
    revision: str = ""
    archivo_origen: str = ""


def generar_reporte(filas: list[FilaReporte], *, titulo: str = "Registros SUNAT") -> bytes:
    """Devuelve el .xlsx (bytes) con una fila por registro."""
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    libro = Workbook(write_only=True)
    hoja = libro.create_sheet(title=titulo[:31] or "Reporte")
    hoja.freeze_panes = "A2"
    for indice, (clave, _) in enumerate(COLUMNAS, start=1):
        hoja.column_dimensions[get_column_letter(indice)].width = _ANCHOS.get(clave, 18)

    negrita = Font(bold=True, color="FFFFFF")
    fondo = PatternFill("solid", fgColor="1F4E79")
    centrado = Alignment(horizontal="center", vertical="center")
    cabecera = []
    for _, etiqueta in COLUMNAS:
        celda = WriteOnlyCell(hoja, value=etiqueta)
        celda.font = negrita
        celda.fill = fondo
        celda.alignment = centrado
        cabecera.append(celda)
    hoja.append(cabecera)

    for fila in filas:
        hoja.append(_a_celdas(hoja, fila))

    bufer = io.BytesIO()
    libro.save(bufer)
    return bufer.getvalue()


def _a_celdas(hoja: object, fila: FilaReporte) -> list[object]:
    from openpyxl.cell import WriteOnlyCell

    celdas: list[object] = []
    for clave, _ in COLUMNAS:
        valor = getattr(fila, clave)
        celda = WriteOnlyCell(hoja, value=valor)
        if clave in _COLUMNAS_NUMERICAS and valor:
            try:
                celda.value = float(Decimal(valor))
                celda.number_format = _FORMATO_IMPORTE
            except (InvalidOperation, ValueError):
                # Importe ilegible: se deja el texto tal cual en lugar de un cero engañoso.
                celda.value = valor
        celdas.append(celda)
    return celdas


def generar_csv(filas: list[FilaReporte]) -> bytes:
    """Devuelve el CSV (bytes, con BOM UTF-8) con una fila por registro."""
    bufer = io.StringIO(newline="")
    escritor = csv.writer(bufer, delimiter=";", quoting=csv.QUOTE_MINIMAL)
    escritor.writerow([etiqueta for _, etiqueta in COLUMNAS])
    for fila in filas:
        escritor.writerow([getattr(fila, clave) for clave, _ in COLUMNAS])
    return bufer.getvalue().encode("utf-8-sig")
